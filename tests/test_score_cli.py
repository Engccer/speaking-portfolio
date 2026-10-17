import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

import score

DATA_PATH = Path(__file__).resolve().parents[1] / "data" / "dialogues.json"
pytestmark = pytest.mark.skipif(not DATA_PATH.exists(), reason="data/dialogues.json 없음: scripts/build_dialogues.py 실행 필요")


def make_student(folder, cls, num, name, ids, perfect=True, practice=False):
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    cdir = folder / ("practice" if practice else "") / f"3-{cls:02d}"; cdir.mkdir(parents=True, exist_ok=True)
    for i, did in enumerate(ids, 1):
        d = next(x for x in data if x["id"] == did)
        idx = "2026-10-09T01-02-03-000Z" if practice else i
        stem = f"3-{cls:02d}-{num:02d} {name}_{idx}_{did}"
        offs, words, t = [], [], 0.0
        for u in d["units"]:
            offs.append({"turn": u["n"], "start": t, "end": t + 10})
            if perfect:
                tt = t
                for tok in u["text"].split():
                    words.append({"text": tok, "start": tt, "end": tt + 0.3, "type": "word"}); tt += 0.4
            t += 10.5
        (cdir / f"{stem}.wav").write_bytes(b"RIFF")
        (cdir / f"{stem}.json").write_text(json.dumps({"dialogue_id": did, "turn_offsets": offs, "duration": t, "submitted_at": "2026-10-12T01:00:00Z"}), encoding="utf-8")
        (cdir / f"{stem}.scribe.json").write_text(json.dumps({"text": "", "words": words}), encoding="utf-8")


def test_score_cli_writes_xlsx_and_csv(tmp_path):
    make_student(tmp_path, 4, 1, "가나다", ["L5-1", "L6-2"])
    make_student(tmp_path, 4, 2, "라마바", ["L5-3", "L7-1"], perfect=False)
    (tmp_path / "미제출_3-04.txt").write_text("03 사아자\n", encoding="utf-8")
    score.run(tmp_path, tmp_path, score.DEFAULT_CFG)
    wb = load_workbook(tmp_path / "4반 말하기 채점.xlsx")
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0][:3] == ("번호", "이름", "대화1")
    by_num = {r[0]: r for r in rows[1:]}
    assert by_num[1][8] == 10            # 최종
    assert by_num[2][8] == 1             # 무음 2편 → 1, 1 → 1
    assert by_num[3][8] == 0 and by_num[3][9] == "미제출"
    csv_text = (tmp_path / "표현별 상세.csv").read_text(encoding="utf-8-sig")
    assert csv_text.count("\n") == 1 + 40  # 헤더 + 2명 × 2편 × 10표현


def test_practice_goes_to_csv_only(tmp_path):
    make_student(tmp_path, 4, 5, "모의만", ["L6-2"], practice=True)
    score.run(tmp_path, tmp_path, score.DEFAULT_CFG)
    assert not (tmp_path / "4반 말하기 채점.xlsx").exists()
    rows = (tmp_path / "표현별 상세.csv").read_text(encoding="utf-8-sig").splitlines()
    assert len(rows) == 11 and rows[1].split(",")[3] == "모의"


def test_returned_submission_is_missing_without_stale_score(tmp_path):
    make_student(tmp_path, 4, 1, "가나다", ["L5-1", "L6-2"])
    for metadata in (tmp_path / "3-04").glob("*.json"):
        if not metadata.name.endswith(".scribe.json"):
            info = json.loads(metadata.read_text(encoding="utf-8"))
            info["returned"] = True
            metadata.write_text(json.dumps(info), encoding="utf-8")
    (tmp_path / "미제출_3-04.txt").write_text("01 가나다", encoding="utf-8")
    score.run(tmp_path, tmp_path, score.DEFAULT_CFG)
    rows = list(load_workbook(tmp_path / "4반 말하기 채점.xlsx").active.iter_rows(values_only=True))
    assert len(rows) == 2
    assert rows[1][0] == 1 and rows[1][8] == 0 and rows[1][9] == "미제출"
    assert len((tmp_path / "표현별 상세.csv").read_text(encoding="utf-8-sig").splitlines()) == 1



def _slow_short_units(scribe_path, max_words, spacing=2.4):
    """단어 수가 max_words 이하인 단위의 단어 간격을 spacing초로 벌린다(단위 창 10초 안에 든다)."""
    data = json.loads(scribe_path.read_text(encoding="utf-8"))
    by_unit = {}
    for w in data["words"]:
        by_unit.setdefault(int(w["start"] // 10.5), []).append(w)
    for words in by_unit.values():
        if len(words) <= max_words:
            base = words[0]["start"]
            for k, w in enumerate(words):
                w["start"], w["end"] = base + k * spacing, base + k * spacing + 0.2
    scribe_path.write_text(json.dumps(data), encoding="utf-8")


def test_note_config_adds_generous_score_to_remark(tmp_path):
    # 간격 2.4초는 기본(2.0초)에선 불만족, 후한 기준(2.5초)에선 만족이다.
    generous = {"match_ratio_min": 0.6, "max_gap_sec": 2.5, "wpm_min": 10}
    make_student(tmp_path, 4, 5, "차카타", ["L5-1", "L6-1"])
    make_student(tmp_path, 4, 6, "파하가", ["L5-1", "L6-1"])
    _slow_short_units(next(tmp_path.glob("3-04/3-04-05 *_1_L5-1.scribe.json")), max_words=0)  # 변경 없음
    _slow_short_units(next(tmp_path.glob("3-04/3-04-06 *_1_L5-1.scribe.json")), max_words=5)  # 짧은 단위 여러 개
    score.run(tmp_path, tmp_path, score.DEFAULT_CFG, generous)
    rows = list(load_workbook(tmp_path / "4반 말하기 채점.xlsx").active.iter_rows(values_only=True))
    hdr = rows[0]; by_num = {r[0]: dict(zip(hdr, r)) for r in rows[1:]}
    assert by_num[5]["최종"] == 10 and "후한 기준" not in (by_num[5]["비고"] or "")
    assert by_num[6]["최종"] < 10 and "후한 기준 적용 시 10점" in by_num[6]["비고"]
