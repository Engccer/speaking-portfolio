"""학생 확인용 인쇄 결과표의 값을 검증하고 게시한다. 기본 실행은 검증만 한다."""
import argparse
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
import re

import openpyxl
from common import service_client


def norm_name(value):
    return re.sub(r"[\s　]", "", value)


def integer(value, label, nullable=False):
    if value is None and nullable:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value or not 0 <= value <= 10:
        raise ValueError(f"{label}: 0~10 정수가 필요합니다.")
    return int(value)


def read_scores(folder):
    records = []
    seen = set()
    paths = sorted(Path(folder).glob("*.xlsx"))
    paths = [p for p in paths if not p.name.startswith("~$")]
    if not paths:
        raise ValueError("인쇄 결과표 xlsx 파일이 없습니다.")
    for path in paths:
        match = re.fullmatch(r"3학년 ([1-6])반 말하기 포트폴리오 결과표.xlsx", path.name)
        if not match:
            raise ValueError(f"예상하지 않은 파일명: {path.name}")
        cls = int(match[1])
        workbook = openpyxl.load_workbook(path, data_only=False)
        try:
            ws = workbook.active
            if ws.title != f"3-{cls}" or ws['A1'].value != f"3학년 {cls}반 말하기 포트폴리오 수행평가 결과표":
                raise ValueError(f"{path.name}: 반 정보 불일치")
            expected = {'A2': '번호', 'B2': '이름', 'C2': '첫 번째 대화', 'G2': '두 번째 대화',
                        'K2': '최종 점수\n(10점 만점)', 'L2': '비고',
                        'C3': '대화문', 'D3': '통과한 표현\n(10개 중)', 'E3': '기준에 못 미친\n표현 번호', 'F3': '점수',
                        'G3': '대화문', 'H3': '통과한 표현\n(10개 중)', 'I3': '기준에 못 미친\n표현 번호', 'J3': '점수'}
            if any(ws[cell].value != value for cell, value in expected.items()):
                raise ValueError(f"{path.name}: 결과표 열 구조 불일치")
            count = 0
            for row in ws.iter_rows(min_row=4, max_col=12):
                values = [c.value for c in row]
                num, name = values[:2]
                if all(v is None for v in values[1:]):
                    continue  # 아래 설명 문단과 빈 줄
                if any(c.data_type == 'f' for c in row):
                    raise ValueError(f"{cls}반 {row[0].row}행: 계산되지 않은 수식은 게시할 수 없습니다.")
                if type(num) is not int or not 1 <= num <= 40 or not isinstance(name, str) or not norm_name(name):
                    raise ValueError(f"{cls}반 {row[0].row}행: 학생 식별 정보 오류")
                key = (cls, num)
                if key in seen:
                    raise ValueError(f"{cls}반 {num}번 중복")
                seen.add(key)
                dialogues = []
                for offset in (2, 6):
                    title, passed, missed, score = values[offset:offset + 4]
                    if all(v is None for v in (title, passed, missed, score)):
                        continue
                    if not isinstance(title, str) or not re.fullmatch(r"Lesson [5-7]-[1-3]", title):
                        raise ValueError(f"{cls}반 {num}번: 대화문 정보 오류")
                    if missed is not None and not re.fullmatch(r"\d+(?:,\s*\d+)*", str(missed)):
                        raise ValueError(f"{cls}반 {num}번: 표현 번호 오류")
                    if missed is not None and any(not 1 <= int(n.strip()) <= 10 for n in str(missed).split(',')):
                        raise ValueError(f"{cls}반 {num}번: 표현 번호 범위 오류")
                    dialogues.append({'title': title, 'passed': integer(passed, '통과한 표현'),
                                      'missed': str(missed) if missed is not None else '', 'score': integer(score, '대화 점수')})
                final = integer(values[10], '최종 점수', nullable=True)
                note = values[11] or ''
                if not isinstance(note, str) or (final is None and not note) or (final is not None and len(dialogues) != 2):
                    raise ValueError(f"{cls}반 {num}번: 점수 또는 상태 정보가 불완전합니다.")
                records.append({'class': cls, 'number': num, 'name': norm_name(name),
                                'final_score': final, 'dialogues': dialogues, 'note': note})
                count += 1
            if not count:
                raise ValueError(f"{path.name}: 학생 행이 없습니다.")
        finally:
            workbook.close()
    return records


def validate_roster(records, roster):
    students = {(s['class'], s['number']): s for s in roster}
    for row in records:
        student = students.get((row['class'], row['number']))
        if not student or student.get('is_teacher') or norm_name(student['name']) != row['name']:
            raise ValueError(f"{row['class']}반 {row['number']}번: 서버 명렬과 불일치")


def payload(records):
    now = datetime.now(timezone.utc).isoformat()
    return [{**{k: v for k, v in r.items() if k != 'name'}, 'published_at': now} for r in records]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('folder', type=Path)
    parser.add_argument('--publish', action='store_true', help='검증 후 서버에 게시한다')
    args = parser.parse_args()
    records = read_scores(args.folder)
    client = service_client()
    validate_roster(records, client.table('students').select('class,number,name,is_teacher').execute().data)
    for cls, count in sorted(Counter(r['class'] for r in records).items()):
        print(f'{cls}반 {count}명 검증 완료')
    if not args.publish:
        print('검증만 완료. 게시하려면 --publish를 지정하세요.')
        return
    rows = payload(records)
    client.table('published_scores').upsert(rows, on_conflict='class,number').execute()
    stored = client.table('published_scores').select('*').execute().data
    lookup = {(r['class'], r['number']): r for r in stored}
    for row in rows:
        actual = lookup.get((row['class'], row['number']), {})
        if any(actual.get(k) != v for k, v in row.items() if k != 'published_at'):
            raise RuntimeError('게시 후 원본 대조 실패')
    print(f'{len(rows)}명 게시 및 원본 전수 대조 완료')


if __name__ == '__main__':
    main()
